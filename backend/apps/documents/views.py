import os
from django.conf import settings
from rest_framework import permissions, status
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import ClientFile, DocumentTree
from .serializers import ClientFileSerializer
from apps.clients.views import get_client_queryset
from rest_framework import generics


class DocumentUploadView(APIView):
    """POST /api/v1/documents/upload/
    Form fields: client (UUID), file (binary), meeting (UUID, optional)"""
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        client_id = request.data.get("client")
        file_obj = request.FILES.get("file")
        meeting_id = request.data.get("meeting")

        if not client_id or not file_obj:
            return Response({"detail": "client and file are required."}, status=status.HTTP_400_BAD_REQUEST)

        qs = get_client_queryset(request.user)
        client = generics.get_object_or_404(qs, pk=client_id)

        # Determine file type from extension
        ext = os.path.splitext(file_obj.name)[1].lower()
        type_map = {
            ".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx",
            ".txt": "txt", ".vtt": "transcript", ".srt": "transcript",
        }
        file_type = type_map.get(ext, "other")

        # Save to media/clients/<client_id>/
        rel_dir = f"clients/{client_id}"
        abs_dir = os.path.join(settings.MEDIA_ROOT, rel_dir)
        os.makedirs(abs_dir, exist_ok=True)
        abs_path = os.path.join(abs_dir, file_obj.name)
        with open(abs_path, "wb") as f:
            for chunk in file_obj.chunks():
                f.write(chunk)

        size_kb = int(file_obj.size / 1024) or 1

        cf = ClientFile.objects.create(
            client=client,
            meeting_id=meeting_id or None,
            name=file_obj.name,
            file_path=f"{rel_dir}/{file_obj.name}",
            file_type=file_type,
            size_kb=size_kb,
            uploaded_by=request.user.username,
        )

        # Queue PageIndex tree build
        from apps.documents.tasks import build_page_index
        build_page_index.delay(str(cf.id))

        return Response(ClientFileSerializer(cf).data, status=status.HTTP_201_CREATED)


class DocumentDeleteView(APIView):
    """DELETE /api/v1/documents/{id}/"""
    permission_classes = [permissions.IsAuthenticated]

    def delete(self, request, pk):
        try:
            cf = ClientFile.objects.get(pk=pk)
        except ClientFile.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)

        # Check ownership
        qs = get_client_queryset(request.user)
        if not qs.filter(pk=cf.client_id).exists():
            return Response(status=status.HTTP_403_FORBIDDEN)

        # Remove physical file
        abs_path = os.path.join(settings.MEDIA_ROOT, cf.file_path)
        if os.path.exists(abs_path):
            os.remove(abs_path)

        cf.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class DocumentSearchView(APIView):
    """POST /api/v1/documents/search/
    Body: { client_id, question, file_id (optional) }"""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        client_id = request.data.get("client_id")
        question = request.data.get("question", "")
        file_id = request.data.get("file_id")

        if not client_id or not question:
            return Response({"detail": "client_id and question are required."}, status=status.HTTP_400_BAD_REQUEST)

        qs = get_client_queryset(request.user)
        client = generics.get_object_or_404(qs, pk=client_id)

        # Load trees for this client
        tree_qs = DocumentTree.objects.filter(client=client, build_status="complete")
        if file_id:
            tree_qs = tree_qs.filter(file_id=file_id)

        if not tree_qs.exists():
            return Response({"detail": "No indexed documents found for this client."}, status=status.HTTP_404_NOT_FOUND)

        # Phase 1: reason over tree to find relevant node, then fetch pages and answer
        from apps.agents.provider import AIProvider

        all_trees = [{"file": str(t.file_id), "name": t.file.name, "tree": t.tree_json} for t in tree_qs]

        tree_str = str(all_trees)[:8000]  # truncate if massive

        node_response = AIProvider().complete(
            system_prompt=(
                "You are a document navigation assistant. Given a JSON tree of document sections, "
                "identify which node_id and which document (file name) best answers the question. "
                "Respond with JSON only: {\"file_name\": \"...\", \"node_id\": \"...\", \"start_index\": N, \"end_index\": N}"
            ),
            user_prompt=f"Document trees:\n{tree_str}\n\nQuestion: {question}",
        )

        import json, re
        match = re.search(r'\{.*\}', node_response, re.DOTALL)
        if not match:
            return Response({"answer": node_response, "source": None})

        try:
            nav = json.loads(match.group())
        except json.JSONDecodeError:
            return Response({"answer": node_response, "source": None})

        # Find matching tree and extract text pages
        target_tree = tree_qs.filter(file__name=nav.get("file_name")).first() or tree_qs.first()
        start = nav.get("start_index", 0)
        end = nav.get("end_index", start + 5)

        abs_path = os.path.join(settings.MEDIA_ROOT, target_tree.file.file_path)
        page_text = _extract_pages(abs_path, target_tree.file.file_type, start, end)

        answer = AIProvider().complete(
            system_prompt="You are a knowledgeable financial advisory assistant. Answer based only on the provided document excerpt.",
            user_prompt=f"Document excerpt (pages {start}–{end}):\n{page_text}\n\nQuestion: {question}",
        )

        return Response({
            "answer": answer,
            "source": {
                "file_name": target_tree.file.name,
                "node_id": nav.get("node_id"),
                "pages": f"{start}–{end}",
            },
        })


def _extract_pages(abs_path, file_type, start, end):
    """Extract text from specific page range of a file."""
    try:
        if file_type in ("pdf", "pdf_image"):
            import pdfplumber
            pages = []
            with pdfplumber.open(abs_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    if start <= i <= end:
                        text = page.extract_text() or ""
                        if not text.strip() and file_type == "pdf_image":
                            import pytesseract
                            from PIL import Image
                            img = page.to_image(resolution=150).original
                            text = pytesseract.image_to_string(img)
                        pages.append(text)
            return "\n\n".join(pages)
        elif file_type == "docx":
            from docx import Document
            doc = Document(abs_path)
            paras = doc.paragraphs[start:end + 1]
            return "\n".join(p.text for p in paras)
        elif file_type == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(abs_path, read_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in list(sheet.iter_rows(values_only=True))[start:end + 1]:
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
        else:
            with open(abs_path, encoding="utf-8", errors="ignore") as f:
                all_lines = f.readlines()
            return "".join(all_lines[start:end + 1])
    except Exception as e:
        return f"[Could not extract text: {e}]"
